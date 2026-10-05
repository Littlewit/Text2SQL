"""导出服务（FR-VIS-20/23）：Excel 导出，复用已脱敏的结果样例。

数据来源是 query_history.result_sample——已在查询链路服务端出口脱敏（FR-SEC-22），
因此导出文件天然无明文敏感值；导出行为本身记录审计（FR-VIS-23）。
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.errors import AppError
from app.infra.models import QueryHistory
from app.services import audit_service


async def export_query_excel(db: AsyncSession, history: QueryHistory, operator_id: int) -> bytes:
    """把查询结果导出为 Excel：数据表 + 元信息（问题/SQL/口径声明）。"""
    if history is None or history.exec_status != "success":
        raise AppError(40400, "仅成功的查询可导出", 404)
    sample = history.result_sample or {}
    columns: list[str] = sample.get("columns", [])
    rows: list[list] = sample.get("rows", [])
    truncated: bool = sample.get("truncated", False)

    wb = Workbook()
    ws = wb.active
    ws.title = "查询结果"

    # 元信息区
    meta_font = Font(bold=True)
    ws.append(["查询问题", history.question])
    ws.append(["生成时间", datetime.now().isoformat(timespec="seconds")])
    ws.append(["口径假设", "；".join((history.assumptions or {}).get("list", [])) or "无"])
    if history.explain_text:
        ws.append(["SQL 解释", history.explain_text])
    ws.append([])

    # 数据表区
    ws.append(columns)
    for cell in ws[ws.max_row]:
        cell.font = meta_font
    for row in rows:
        ws.append(row)
    if truncated:
        ws.append([])  # type: ignore[arg-type] —— 空行分隔
        ws.append([f"（结果已按行数上限截断，仅含前 {len(rows)} 行）"])

    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()

    await audit_service.record(
        db, user_id=operator_id, action="export.excel",
        object_type="query_history", object_id=str(history.id),
        detail={"rows": len(rows), "truncated": truncated},
    )
    await db.commit()
    return content


def suggest_followups(question: str, chart_type: str | None) -> list[str]:
    """建议追问（FR-UI-04）：基于问题与图表类型的规则推荐。"""
    suggestions: list[str] = ["按月拆分看趋势", "换成表格展示", "导出 Excel"]
    if any(k in question for k in ("哪个", "最高", "最低", "排名", "Top")):
        suggestions.insert(0, "看前 10 名的完整排名")
    if chart_type == "bar":
        suggestions.insert(0, "换成分组柱状图对比")
    if "退货" in question:
        suggestions.insert(0, "只看退货率超过 5% 的")
    return suggestions[:4]
